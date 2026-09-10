"""Cover letter assembly, with the same gate the resume has.

The Resume Engine's rule -- selection among approved statements, never
invention -- is worth more here than it is on the resume, because a cover
letter is free prose and free prose is where a claim gets invented. So the
letter is built the same way: paragraphs come out of `content/letter.toml`
unchanged, only declared slots are filled, and the result is checked before
it is written.

Three checks, in `verify`:

1. **Template fidelity.** Each rendered paragraph must match its own template
   with only the declared slots substituted. Nothing can be appended, edited,
   or smuggled in.
2. **No unsupported numbers.** Every number in the letter must also appear in
   the tailored resume shipping with it. This is the "a variant may reword a
   claim, never renumber it" rule ported across -- and it has teeth here,
   because it drops a paragraph whose evidence the resume didn't select.
3. **No unsupported tools.** Every tool named must be on the resume's SKILLS
   line and in the job description.

The only text exempt from check 1 is a company note you type yourself,
which is quoted back to you for review and recorded as your own words.
"""

from __future__ import annotations

import re
import tomllib
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from . import config

_SLOT = re.compile(r"\{(\w+)\}")
_NUMBER = re.compile(r"\b\d[\d,]*(?:\.\d+)?\b")
_WORD_NUMBERS = {
    "one": "1", "two": "2", "three": "3", "four": "4", "five": "5",
    "six": "6", "seven": "7", "eight": "8", "nine": "9", "ten": "10",
    "eleven": "11", "twelve": "12",
}
_WORD_NUMBER_RE = re.compile(r"\b(" + "|".join(_WORD_NUMBERS) + r")\b", re.I)


@dataclass
class Paragraph:
    id: str
    template: str
    text: str
    slots: dict[str, str] = field(default_factory=dict)
    user_written: bool = False


@dataclass
class Letter:
    company: str
    role: str
    header: list[str]
    greeting: str
    paragraphs: list[Paragraph]
    sign_off: str
    name: str
    tools: list[str] = field(default_factory=list)
    dropped: list[tuple[str, str]] = field(default_factory=list)

    @property
    def body(self) -> str:
        return "\n\n".join(p.text for p in self.paragraphs)

    def to_text(self) -> str:
        lines = list(self.header)
        lines += ["", date.today().strftime("%B %d, %Y"), "",
                  f"{self.company}", "", self.greeting, "", self.body, "",
                  self.sign_off, self.name]
        return "\n".join(lines).strip() + "\n"

    def to_markdown(self) -> str:
        head = " | ".join(h for h in self.header if h)
        out = [f"# Cover letter -- {self.company}, {self.role}", "",
               f"*{head}*", "", f"*{date.today().strftime('%B %d, %Y')}*", "",
               self.greeting, ""]
        out += [p.text for p in self.paragraphs]
        out += ["", self.sign_off, "", self.name]
        if self.dropped:
            out += ["", "---", "",
                    "## Paragraphs the gate dropped", "",
                    "These are approved paragraphs that were *not* used, "
                    "because the tailored resume shipping with this letter "
                    "does not make the claim they rest on. That is the check "
                    "working, not a bug.", ""]
            out += [f"- `{pid}` -- {why}" for pid, why in self.dropped]
        return "\n".join(out) + "\n"


# ---------------------------------------------------------------------------

def load_content(path: Path | None = None) -> dict:
    return tomllib.loads((path or config.LETTER_FILE).read_text(encoding="utf-8"))


def numbers_in(text: str) -> set[str]:
    """Every number in a piece of text, normalized for comparison.

    Commas are stripped so "2,000" and "2000" are the same claim, and spelled
    numbers are folded in so "four departments" is checked against "4".
    """
    found = {m.group(0).replace(",", "").rstrip(".").lstrip("0") or "0"
             for m in _NUMBER.finditer(text)}
    found |= {_WORD_NUMBERS[m.group(1).lower()]
              for m in _WORD_NUMBER_RE.finditer(text)}
    return found


def resume_contact(resume_text: str) -> tuple[str, list[str]]:
    """Name and contact line, lifted from the tailored resume itself.

    Deliberately not duplicated into letter.toml: two copies of a phone
    number is two chances to send the wrong one.
    """
    lines = [ln.strip() for ln in resume_text.splitlines() if ln.strip()]
    if not lines:
        return "", []
    name = lines[0]
    contact = lines[1] if len(lines) > 1 else ""
    return name, [name, contact]


def resume_skills(resume_text: str) -> list[str]:
    """Every skill named on the resume's SKILLS section."""
    skills: list[str] = []
    in_section = False
    for line in resume_text.splitlines():
        stripped = line.strip()
        if stripped == "SKILLS":
            in_section = True
            continue
        if in_section:
            if not stripped:
                continue
            if stripped.isupper() and ":" not in stripped:
                break
            _, _, rest = stripped.partition(":")
            for item in (rest or stripped).split(","):
                item = item.strip()
                # "Excel (Pivot Tables, VLOOKUP)" splits on the comma inside
                # the parenthetical; keep the head, drop the fragments.
                if not item or item.endswith(")") and "(" not in item:
                    continue
                skills.append(item.split("(")[0].strip())
    return [s for s in skills if s]


def pick_tools(resume_text: str, jd_text: str, limit: int = 3,
               stoplist: list[str] | None = None) -> list[str]:
    """Tools the resume claims AND the JD asked for, most-wanted first.

    `stoplist` drops skills that are true and worth a keyword match but read
    as filler in a sentence -- "what I have kept working with is Microsoft
    Office" argues against itself.
    """
    low_jd = jd_text.lower()
    blocked = {s.lower() for s in (stoplist or [])}
    scored: list[tuple[int, str]] = []
    for skill in resume_skills(resume_text):
        needle = skill.lower()
        if needle in blocked:
            continue
        count = low_jd.count(needle)
        if count and len(needle) > 2:
            scored.append((count, skill))
    scored.sort(key=lambda pair: (-pair[0], pair[1]))
    picked: list[str] = []
    for _, skill in scored:
        if not any(skill.lower() in p.lower() or p.lower() in skill.lower()
                   for p in picked):
            picked.append(skill)
        if len(picked) == limit:
            break
    return picked


def join_tools(tools: list[str]) -> str:
    # Empty is a real case, not a bug: a JD can name no tool this resume also
    # names, and `requires_tools` exists precisely so the paragraph that needs
    # one steps aside. The slot still gets computed before that check runs, so
    # this has to answer rather than raise -- it crashed the whole build on the
    # first JD with no overlap (Renewals Analyst, 2026-08-14).
    if not tools:
        return ""
    if len(tools) == 1:
        return tools[0]
    if len(tools) == 2:
        return f"{tools[0]} and {tools[1]}"
    return ", ".join(tools[:-1]) + f", and {tools[-1]}"


def _matches_family(entry: dict, family: str) -> bool:
    families = entry.get("families") or ["*"]
    return "*" in families or family in families


def _render(entry: dict, slots: dict[str, str]) -> Paragraph:
    template = entry["template"].strip()
    used = {name: slots.get(name, "") for name in _SLOT.findall(template)}
    text = template
    for name, value in used.items():
        text = text.replace("{" + name + "}", value)
    return Paragraph(id=entry["id"], template=template, text=text, slots=used)


def _supported(paragraph: Paragraph, resume_numbers: set[str]) -> str:
    """"" if the paragraph's numbers are all on the resume, else why not."""
    missing = sorted(numbers_in(paragraph.text) - resume_numbers)
    # A slot value is your own text or a tool name; its numbers are not
    # claims the paragraph is making.
    for value in paragraph.slots.values():
        missing = [n for n in missing if n not in numbers_in(value)]
    if missing:
        return (f"the resume does not state {', '.join(missing)}, so the "
                f"letter cannot either")
    return ""


def build(*, company: str, role: str, family: str, resume_text: str,
          jd_text: str, note: str = "", content: dict | None = None,
          evidence_count: int = 2) -> Letter:
    content = content or load_content()
    meta = content.get("meta", {})
    resume_numbers = numbers_in(resume_text)
    tools = pick_tools(resume_text, jd_text,
                       stoplist=meta.get("tool_stoplist"))
    slots = {"company": company, "role": role, "tools": join_tools(tools)}

    dropped: list[tuple[str, str]] = []
    chosen: list[Paragraph] = []

    def take(section: str, count: int = 1, **extra) -> None:
        entries = [e for e in content.get(section, [])
                   if _matches_family(e, family)]
        # Family-specific entries beat the "*" catch-all.
        entries.sort(key=lambda e: "*" in (e.get("families") or ["*"]))
        taken = 0
        # Two paragraphs that both open "At Spargo I..." read as one story told
        # twice. Where entries declare an `employer`, spend each one once --
        # but a repeated employer is better than a short letter, so the pass
        # runs twice and the second one stops caring.
        spent: set[str] = set()
        seen: set[str] = set()
        deferred: list[dict] = []

        def consider(entry: dict, *, unique: bool) -> None:
            nonlocal taken
            if taken >= count or entry["id"] in seen:
                return
            if "requires_tools" in entry and entry["requires_tools"] != bool(tools):
                return
            if extra.get("skip_ids") and entry["id"] in extra["skip_ids"]:
                return
            employer = entry.get("employer")
            if unique and employer and employer in spent:
                deferred.append(entry)
                return
            paragraph = _render(entry, slots)
            seen.add(entry["id"])
            reason = _supported(paragraph, resume_numbers)
            if reason:
                dropped.append((entry["id"], reason))
                return
            chosen.append(paragraph)
            if employer:
                spent.add(employer)
            taken += 1

        for entry in entries:
            consider(entry, unique=True)
        for entry in deferred:
            consider(entry, unique=False)

    take("opening", 1)
    take("evidence", evidence_count)
    if note.strip():
        # The one exemption: your own sentence about the company. Quoted
        # back in the packet so it is reviewed, and marked as your words.
        chosen.append(Paragraph(id="note.user", template="", text=note.strip(),
                                user_written=True))
    take("bridge", 1)
    take("close", 1)

    name, header = resume_contact(resume_text)
    return Letter(
        company=company, role=role, header=header,
        greeting=meta.get("greeting", "Dear Hiring Team,"),
        paragraphs=chosen, sign_off=meta.get("sign_off", "Sincerely,"),
        name=name, tools=tools, dropped=dropped,
    )


def verify(letter: Letter, resume_text: str, jd_text: str) -> list[str]:
    """The gate. A non-empty return means do not deliver this letter."""
    problems: list[str] = []
    resume_numbers = numbers_in(resume_text)
    low_resume, low_jd = resume_text.lower(), jd_text.lower()

    for paragraph in letter.paragraphs:
        if paragraph.user_written:
            continue
        # 1. template fidelity -- rebuild the regex from the template so any
        #    edit to the rendered text fails, not just an obvious one.
        pattern = re.escape(paragraph.template)
        for name in _SLOT.findall(paragraph.template):
            pattern = pattern.replace(re.escape("{" + name + "}"),
                                      re.escape(paragraph.slots.get(name, "")))
        if not re.fullmatch(pattern, paragraph.text, re.S):
            problems.append(f"{paragraph.id}: rendered text does not match its "
                            f"approved template")

        # 2. numbers
        slot_numbers: set[str] = set()
        for value in paragraph.slots.values():
            slot_numbers |= numbers_in(value)
        unsupported = numbers_in(paragraph.text) - resume_numbers - slot_numbers
        if unsupported:
            problems.append(f"{paragraph.id}: claims {', '.join(sorted(unsupported))}, "
                            f"which the tailored resume does not say")

    # 3. tools
    for tool in letter.tools:
        if tool.lower() not in low_resume:
            problems.append(f"tool '{tool}' is named in the letter but is not "
                            f"on the resume")
        if tool.lower() not in low_jd:
            problems.append(f"tool '{tool}' is named in the letter but the job "
                            f"description never asked for it")

    if not any(p.id.startswith("open") for p in letter.paragraphs):
        problems.append("no opening paragraph survived selection")
    if not any(p.id.startswith("close") for p in letter.paragraphs):
        problems.append("no closing paragraph survived selection")
    return problems


_ASCII = {
    "‘": "'", "’": "'", "“": '"', "”": '"',
    "–": "-", "—": "-", "…": "...", " ": " ",
    "•": "-", "·": "-", "−": "-",
}


def ascii_safe(text: str) -> str:
    """Core PDF fonts are Latin-1; anything outside it raises on render.

    Same trade the Resume Engine makes: a document full of typographic quotes
    parses worse anyway, so ASCII is the safe surface.
    """
    for bad, good in _ASCII.items():
        text = text.replace(bad, good)
    return text.encode("latin-1", "replace").decode("latin-1")


def normalize_pdf_text(text: str) -> str:
    """Collapse whitespace so PDF line-wrapping doesn't break comparison."""
    return re.sub(r"\s+", " ", ascii_safe(text)).strip()


def scrub_docx(doc, name: str) -> None:
    """Strip the toolchain's fingerprints off a .docx before it ships.

    python-docx builds every document from its own default template, and that
    template arrives pre-filled: creator "python-docx", description "generated
    by python-docx", and a created/modified date in December 2013 that is the
    day the template was authored. None of that is visible in the document, all
    of it is one right-click away in Windows' file properties, and together it
    says "a script wrote this" to anyone who looks.

    So the properties get set to what is actually true: you wrote it, today.
    """
    from datetime import datetime

    props = doc.core_properties
    props.author = name
    props.last_modified_by = name
    props.title = ""
    props.subject = ""
    props.comments = ""          # the "generated by python-docx" line
    props.category = ""
    props.keywords = ""
    props.created = props.modified = datetime.now()
    props.revision = 1

    # app.xml carries the template's own <Application>Microsoft Macintosh
    # Word</Application>, which python-docx does not expose. Claiming to be a
    # Word build that never touched this file is a worse lie than saying
    # nothing, so the element is emptied rather than replaced.
    try:
        part = next(p for p in doc.part.package.parts
                    if p.partname == "/docProps/app.xml")
        xml = part.blob.decode("utf-8")
        xml = re.sub(r"<Application>.*?</Application>", "<Application></Application>", xml)
        xml = re.sub(r"<Company>.*?</Company>", "<Company></Company>", xml)
        xml = re.sub(r"<Template>.*?</Template>", "<Template></Template>", xml)
        part._blob = xml.encode("utf-8")
    except (StopIteration, AttributeError):
        pass


def write_docx(letter: Letter, path: Path) -> Path | None:
    """A .docx copy for the portals that parse Word more reliably.

    Both formats ship for the same reason the resume ships both: Workday and
    Taleo read .docx better, most other systems are happy with the PDF, and
    generating both removes the guess. The PDF is the one to attach when the
    portal doesn't care.

    No tables and no header/footer, for the same reason the Resume Engine
    avoids them: they are the classic ATS parse trap.
    """
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        return None
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for line in letter.header:
        doc.add_paragraph(line)
    doc.add_paragraph("")
    doc.add_paragraph(date.today().strftime("%B %d, %Y"))
    doc.add_paragraph(letter.company)
    doc.add_paragraph("")
    doc.add_paragraph(letter.greeting)
    for paragraph in letter.paragraphs:
        doc.add_paragraph("")
        doc.add_paragraph(paragraph.text)
    doc.add_paragraph("")
    doc.add_paragraph(letter.sign_off)
    doc.add_paragraph(letter.name)
    scrub_docx(doc, letter.name)
    doc.save(str(path))
    return path


def write_pdf(letter: Letter, path: Path) -> Path | None:
    """The letter as a PDF -- the copy that gets uploaded. Best-effort.

    Same layout rules as the resume: one column, one text flow, a core
    PostScript font, and no tables, text boxes, headers or footers. A cover
    letter is read by a person more often than by a parser, but it goes up the
    same pipe as the resume and gets read by the same machine first.
    """
    try:
        from fpdf import FPDF
    except ImportError:
        return None

    pdf = FPDF("P", "mm", "Letter")
    pdf.set_auto_page_break(auto=True, margin=18)
    # Same document properties the Resume Engine ships: empty rather than
    # advertising the library. fpdf2 fills Producer in by default, and a PDF
    # whose Producer says "py-fpdf" is the one obvious tell in the packet.
    pdf.set_title("")
    pdf.set_author("")
    for setter in ("set_producer", "set_creator", "set_subject", "set_keywords"):
        if hasattr(pdf, setter):
            getattr(pdf, setter)("")
    pdf.add_page()
    pdf.set_margins(22, 20, 22)
    pdf.set_y(20)

    width = pdf.w - pdf.l_margin - pdf.r_margin

    def block(text: str, *, size: float = 10.5, bold: bool = False,
              gray: int = 40, lh: float = 5.0) -> None:
        pdf.set_font("Helvetica", "B" if bold else "", size)
        pdf.set_text_color(gray, gray, gray)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(width, lh, ascii_safe(text))

    if letter.header:
        block(letter.header[0], size=15, bold=True, gray=20, lh=7)
        for line in letter.header[1:]:
            block(line, size=9, gray=80, lh=4.5)
    pdf.ln(6)

    block(date.today().strftime("%B %d, %Y"), gray=80)
    block(letter.company)
    pdf.ln(5)
    block(letter.greeting)

    for paragraph in letter.paragraphs:
        pdf.ln(4)
        block(paragraph.text, lh=5.2)

    pdf.ln(6)
    block(letter.sign_off)
    block(letter.name)

    pdf.output(str(path))
    return path
